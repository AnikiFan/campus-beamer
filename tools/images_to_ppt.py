#!/usr/bin/env python3
# Campus Beamer: render PDF pages with links, notes, sections and metadata.
# Copyright (C) 2026 Xiao Fan <xiaofan140@gmail.com>
# SPDX-License-Identifier: GPL-3.0-or-later
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.
"""Render PDF pages as PPTX images and recreate links and local video media.

Run from the project root: uv run tools/images_to_ppt.py build/main.pdf
Native ``build/main.notes.pdf`` renders produced from TeX are read into the
PowerPoint speaker notes; legacy ``.notes.json`` imports remain supported.
``--notes-dir`` selects a different notes directory. Local video links are embedded as PowerPoint movie
objects with a first-frame poster and proportional contained fit when the file
exists. FFmpeg is required only for these local videos. Other links become
clickable areas, and the original PDF is never modified.
Level-one PDF bookmarks become native PowerPoint sections. Standard PDF
document properties and the Campus class's cover metadata are copied to PPTX.
Only PyMuPDF and python-pptx are direct dependencies. Slide content remains an
image; each PDF page (including each Beamer overlay) becomes one slide.
PDF Fade/Dissolve transitions and explicit page durations become native PPTX
slide transitions; no effect is added to pages without transition settings.
"""

from __future__ import annotations

import argparse
import hashlib
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from io import BytesIO
import json
import math
import os
import re
from pathlib import Path, PureWindowsPath
import shutil
import subprocess
import sys
import tempfile
from urllib.parse import quote, unquote, urlsplit
from uuid import uuid4
from xml.etree import ElementTree as ET

import pymupdf
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import CONTENT_TYPE, RELATIONSHIP_TYPE
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches


@dataclass
class LinkSummary:
    internal: int = 0
    external: int = 0
    embedded: int = 0
    skipped: int = 0


def campus_metadata(pdf) -> dict[str, str]:
    """Read the class's Unicode PDF Info fields without parsing TeX source."""
    kind, reference = pdf.xref_get_key(-1, 'Info')
    if kind != 'xref':
        return {}
    info_xref = int(reference.split()[0])
    fields = {}
    for name in ('Title', 'Subtitle', 'Group', 'Advisor', 'Date', 'Contact', 'Language'):
        kind, value = pdf.xref_get_key(info_xref, 'Campus' + name)
        if kind == 'string' and value.strip():
            fields[name] = value.strip()
    return fields


def pdf_creation_time(value: str) -> datetime | None:
    """Convert a PDF timestamp with an explicit timezone to naive UTC."""
    match = re.fullmatch(r"D:(\d{14})(Z|[+-]\d{2}'?\d{2}'?)", value)
    if not match:
        return None
    try:
        result = datetime.strptime(match[1], '%Y%m%d%H%M%S')
        zone = match[2].replace("'", '')
        offset = timedelta() if zone == 'Z' else timedelta(
            hours=int(zone[1:3]), minutes=int(zone[3:5])) * (1 if zone[0] == '+' else -1)
        return result.replace(tzinfo=timezone(offset)).astimezone(timezone.utc).replace(tzinfo=None)
    except ValueError:
        return None


def set_document_metadata(prs, pdf, campus: dict[str, str]) -> None:
    """Copy document properties, keeping the talk date separate from file dates."""
    metadata = pdf.metadata
    core = prs.core_properties
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    core.created = pdf_creation_time(metadata.get('creationDate', '')) or now
    core.modified = now
    core.last_modified_by = ''
    core.comments = ''
    custom = {name: campus[key] for key, name in (
        ('Subtitle', 'Subtitle'), ('Group', 'Group'), ('Advisor', 'Advisor'),
        ('Date', 'PresentationDate'), ('Contact', 'Contact')) if key in campus}
    for name in ('title', 'author', 'subject', 'keywords'):
        value = campus.get('Title', metadata.get('title', '')) if name == 'title' else metadata.get(name, '')
        value = value or ''
        # Office core properties have a 255-character limit; retain the full
        # value in custom properties rather than silently discarding it.
        if len(value) > 255:
            custom['Full' + name.title()] = value
            print(f'WARNING: {name} exceeds 255 characters; full value stored in custom properties.', file=sys.stderr)
        setattr(core, name, value[:255])
    if not custom:
        return
    namespace = 'http://schemas.openxmlformats.org/officeDocument/2006/custom-properties'
    types = 'http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'
    ET.register_namespace('vt', types)
    root = ET.Element(f'{{{namespace}}}Properties')
    for pid, (name, value) in enumerate(custom.items(), start=2):
        prop = ET.SubElement(root, f'{{{namespace}}}property', {
            'fmtid': '{D5CDD505-2E9C-101B-9397-08002B2CF9AE}', 'pid': str(pid), 'name': name})
        ET.SubElement(prop, f'{{{types}}}lpwstr').text = value
    part = Part(PackURI('/docProps/custom.xml'), CONTENT_TYPE.OFC_CUSTOM_PROPERTIES,
                prs.part.package, ET.tostring(root, encoding='utf-8', xml_declaration=True))
    prs.part.package.relate_to(part, RELATIONSHIP_TYPE.CUSTOM_PROPERTIES)


def set_presentation_sections(prs, pdf, campus: dict[str, str]) -> int:
    """Map level-one PDF bookmarks to flat native PowerPoint sections."""
    starts = []
    for level, title, page in pdf.get_toc():
        if level != 1:
            continue
        if page == -1:
            print(f'WARNING: section bookmark {title!r} has no local page; skipped.', file=sys.stderr)
            continue
        if not 1 <= page <= len(prs.slides):
            raise ValueError(f'Section bookmark {title!r} points outside the PDF: {page}')
        index = page - 1
        if starts and index < starts[-1][1]:
            raise ValueError('Section bookmarks are not in PDF page order. Output was not saved.')
        if starts and index == starts[-1][1]:
            print(f'WARNING: section {starts[-1][0]!r} has no pages; skipped.', file=sys.stderr)
            starts.pop()
        starts.append((title, index))
    if not starts:
        return 0
    if starts[0][1] > 0:
        name = '封面与导读' if campus.get('Language') == 'chinese' else 'Introduction'
        starts.insert(0, (name, 0))
    namespace = 'http://schemas.microsoft.com/office/powerpoint/2010/main'
    ET.register_namespace('p14', namespace)
    sections = ET.Element(f'{{{namespace}}}sectionLst')
    for position, (name, start) in enumerate(starts):
        stop = starts[position + 1][1] if position + 1 < len(starts) else len(prs.slides)
        section = ET.SubElement(sections, f'{{{namespace}}}section', {
            'name': name, 'id': '{' + str(uuid4()).upper() + '}'})
        ids = ET.SubElement(section, f'{{{namespace}}}sldIdLst')
        for index in range(start, stop):
            ET.SubElement(ids, f'{{{namespace}}}sldId', {'id': str(prs.slides[index].slide_id)})
    extensions = prs._element.find(qn('p:extLst'))
    if extensions is None:
        extensions = OxmlElement('p:extLst')
        prs._element.append(extensions)
    extension = OxmlElement('p:ext')
    extension.set('uri', '{521415D9-36F7-43E2-AB2F-B90AF26B5E84}')
    extension.append(parse_xml(ET.tostring(sections, encoding='utf-8')))
    extensions.append(extension)
    return len(starts)


def set_export_properties(prs, video_count: int) -> None:
    """Replace the blank template's stale application name and page counts."""
    namespace = 'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'
    root = ET.Element(f'{{{namespace}}}Properties')
    ratio = prs.slide_width / prs.slide_height
    format_name = 'Widescreen (16:9)' if abs(ratio - 16 / 9) < 0.001 else (
        'On-screen Show (4:3)' if abs(ratio - 4 / 3) < 0.001 else 'Custom')
    for name, value in (('Application', 'Campus Beamer'), ('PresentationFormat', format_name),
                        ('Slides', len(prs.slides)), ('Notes', len(prs.slides)),
                        ('HiddenSlides', 0), ('MMClips', video_count)):
        ET.SubElement(root, f'{{{namespace}}}{name}').text = str(value)
    part = prs.part.package.part_related_by(RELATIONSHIP_TYPE.EXTENDED_PROPERTIES)
    # python-pptx loads app.xml as a generic OPC Part without a public XML setter.
    part._blob = ET.tostring(root, encoding='utf-8', xml_declaration=True)


VIDEO_MIME_TYPES = {
    '.avi': 'video/avi',
    '.m4v': 'video/mp4',
    '.mkv': 'video/x-matroska',
    '.mov': 'video/quicktime',
    '.mp4': 'video/mp4',
    '.mpeg': 'video/mpeg',
    '.mpg': 'video/mpeg',
    '.webm': 'video/webm',
    '.wmv': 'video/x-ms-wmv',
}


def load_notes(notes_path: str | Path | None, page_count: int, strict: bool = False) -> list[str]:
    """Load one speaker-note string per final PDF page.

    Accepted JSON forms are ``["note for page 1", ...]`` and
    ``{"slides": [{"notes": "..."}, ...]}``. Missing trailing entries are
    treated as empty notes; extra entries are rejected to prevent page drift.
    """
    if notes_path is None:
        return [''] * page_count
    path = Path(notes_path)
    if not path.is_file():
        raise FileNotFoundError(path)
    try:
        data = json.loads(path.read_text(encoding='utf-8'))
    except json.JSONDecodeError as exc:
        raise ValueError(f'Invalid notes JSON in {path}: {exc}') from exc

    if isinstance(data, dict):
        data = data.get('slides', data.get('notes'))
    if not isinstance(data, list):
        raise ValueError(f'Notes file {path} must contain a JSON list or a slides list.')

    notes: list[str] = []
    for index, item in enumerate(data, start=1):
        if isinstance(item, str):
            text = item
        elif isinstance(item, dict):
            text = item.get('notes', item.get('speaker_notes', item.get('text', '')))
        else:
            raise ValueError(f'Notes entry {index} in {path} must be a string or object.')
        if not isinstance(text, str):
            raise ValueError(f'Notes entry {index} in {path} must contain text.')
        notes.append(text)
    if len(notes) > page_count:
        raise ValueError(
            f'Notes file {path} has {len(notes)} entries, but the PDF has only {page_count} pages.'
        )
    if strict and len(notes) != page_count:
        raise ValueError(
            f'Notes file {path} has {len(notes)} entries, but the PDF has {page_count} pages. '
            'Provide one entry per page, using an empty string for intentional blanks.'
        )
    return notes + [''] * (page_count - len(notes))


def pdf_info_string(pdf, key: str) -> str:
    """Read a named Info string without inferring layout from page similarity."""
    kind, reference = pdf.xref_get_key(-1, 'Info')
    if kind != 'xref':
        return ''
    kind, value = pdf.xref_get_key(int(reference.split()[0]), key)
    return value if kind == 'string' else ''


def load_native_notes(path: str | Path, slides_pdf, slides_path: Path) -> list[str]:
    """Extract native Beamer note text from an explicitly marked right screen."""
    with pymupdf.open(path) as notes_pdf:
        if not notes_pdf.is_pdf or notes_pdf.needs_pass:
            raise ValueError('Native notes must be an unencrypted PDF.')
        if pdf_info_string(notes_pdf, 'CampusNotesLayout') != 'right':
            raise ValueError('Native notes PDF must use Campus second-screen notes.')
        if len(notes_pdf) != len(slides_pdf):
            raise ValueError('Native notes PDF and slide PDF have different page counts.')
        fingerprint = pdf_info_string(notes_pdf, 'CampusSlidesSHA256')
        if fingerprint and fingerprint != hashlib.sha256(slides_path.read_bytes()).hexdigest():
            raise ValueError('Native notes PDF does not match this slide PDF. Rebuild with make.')
        result = []
        for page, slide in zip(notes_pdf, slides_pdf):
            if (page.rotation or abs(page.rect.width - 2 * slide.rect.width) > .5
                    or abs(page.rect.height - slide.rect.height) > .5):
                raise ValueError('Native notes PDF has an unexpected second-screen geometry.')
            right = pymupdf.Rect(page.rect.width / 2, 0, page.rect.width, page.rect.height)
            result.append(page.get_text('text', clip=right, sort=True).strip())
        return result


def set_slide_notes(slide, text: str) -> None:
    """Write speaker notes through the public python-pptx notes API."""
    frame = slide.notes_slide.notes_text_frame
    frame.clear()
    frame.text = text


def valid_page(page: object, page_count: int) -> bool:
    return type(page) is int and 0 <= page < page_count


def lookup_named_destination(name: str, destinations: dict, page_count: int) -> int | None:
    # Try the literal name first: '%' and '/' can be part of a PDF name.
    for candidate in dict.fromkeys((name, unquote(name), name.removeprefix('/'))):
        page = destinations.get(candidate, {}).get('page')
        if valid_page(page, page_count):
            return page
    return None


def resolve_internal_link(
    link: dict, source_page: int, page_count: int, named_destinations: dict
) -> int | None:
    """Resolve GoTo/named links, without mistaking a remote PDF page for a slide."""
    kind = link.get('kind')
    if kind not in (pymupdf.LINK_GOTO, pymupdf.LINK_NAMED, pymupdf.LINK_URI):
        return None
    if kind != pymupdf.LINK_URI and valid_page(link.get('page'), page_count):
        return link['page']

    candidates = [link.get(key) for key in ('nameddest', 'name', 'named', 'to', 'uri')]
    for candidate in candidates:
        if not isinstance(candidate, str):
            continue
        if kind == pymupdf.LINK_URI and not candidate.startswith('#'):
            continue
        text = candidate.removeprefix('#')
        if text.startswith('nameddest='):
            # Explicit named destinations must not be guessed as page numbers.
            page = lookup_named_destination(text[len('nameddest='):], named_destinations, page_count)
            if page is not None:
                return page
            continue
        page = lookup_named_destination(text, named_destinations, page_count)
        if page is not None:
            return page
        if text.startswith('page=') or (candidate.startswith('#') and text.isdigit()):
            number = text.removeprefix('page=').split('&', 1)[0]
            if number.isdigit() and valid_page(int(number) - 1, page_count):
                return int(number) - 1
        actions = {
            'NextPage': min(source_page + 1, page_count - 1),
            'PrevPage': max(source_page - 1, 0),
            'PreviousPage': max(source_page - 1, 0),
            'FirstPage': 0,
            'LastPage': page_count - 1,
        }
        if text in actions:
            return actions[text]
    return None


def file_link(filename: str, pdf_path: Path, pptx_path: Path) -> str:
    """Rebase relative file links so moving the output directory remains safe."""
    if PureWindowsPath(filename).is_absolute():
        return PureWindowsPath(filename).as_uri()
    path = Path(unquote(filename.replace('\\', '/')))
    if path.is_absolute():
        return path.as_uri()
    target = pdf_path.parent / path
    try:
        return quote(os.path.relpath(target, pptx_path.parent), safe='/')
    except ValueError:  # Windows: input and output on different drives.
        return target.resolve().as_uri()


def resolve_external_link(link: dict, pdf_path: Path, pptx_path: Path) -> str | None:
    kind = link.get('kind')
    if kind == pymupdf.LINK_URI:
        uri = link.get('uri', '')
        if not uri or uri.startswith('#'):
            return None
        parts = urlsplit(uri)
        if parts.scheme.lower() == 'file' and not parts.netloc:
            uri = file_link(unquote(parts.path), pdf_path, pptx_path)
            if parts.fragment:
                uri += '#' + parts.fragment
        return uri
    if kind not in (pymupdf.LINK_GOTOR, pymupdf.LINK_LAUNCH):
        return None
    filename = link.get('file')
    if not filename:
        return None
    uri = file_link(filename, pdf_path, pptx_path)
    # Some TeX producers use GoToR for videos, with a meaningless page=0.
    # Preserve a page/position fragment only for links to another PDF.
    if kind == pymupdf.LINK_GOTOR and filename.lower().endswith('.pdf'):
        page = link.get('page', -1)
        if type(page) is int and page >= 0:
            uri += f'#page={page + 1}'
        elif isinstance(link.get('to'), str):
            uri += '#nameddest=' + quote(link['to'], safe='')
    return uri


def local_file_name(link: dict) -> str | None:
    """Return the file name carried by a local PDF file action.

    A normal web URL is deliberately excluded: it cannot be embedded during
    a deterministic offline conversion. ``file:`` URLs and GoToR/Launch
    actions may both refer to local files.
    """
    kind = link.get('kind')
    if kind == pymupdf.LINK_URI:
        uri = link.get('uri', '')
        if not isinstance(uri, str):
            return None
        parts = urlsplit(uri)
        if parts.scheme.lower() != 'file' or parts.netloc not in ('', 'localhost'):
            return None
        return unquote(parts.path)
    if kind in (pymupdf.LINK_GOTOR, pymupdf.LINK_LAUNCH):
        filename = link.get('file')
        return unquote(filename) if isinstance(filename, str) else None
    return None


def local_video_for_link(link: dict, pdf_path: Path) -> tuple[Path, str] | None:
    """Resolve an existing local video and its MIME type for embedding."""
    filename = local_file_name(link)
    if filename is None:
        return None
    suffix = Path(filename.replace('\\', '/')).suffix.lower()
    mime_type = VIDEO_MIME_TYPES.get(suffix)
    if mime_type is None:
        return None
    path = PureWindowsPath(filename)
    if Path(filename).is_absolute():
        target = Path(filename)
    elif path.is_absolute():
        target = Path(path.as_posix())
    else:
        target = pdf_path.parent / Path(filename.replace('\\', '/'))
    if not target.is_file():
        return None
    return target.resolve(), mime_type


def video_poster(video_path: Path) -> bytes:
    """Decode a real first-frame poster; normalize pixel aspect and rotation.

    FFmpeg is a conditional system dependency, not another Python package.
    Failure is fatal so we never replace a working export with a blank movie.
    """
    executable = shutil.which('ffmpeg')
    if executable is None:
        raise RuntimeError('Embedding local videos requires FFmpeg. Install ffmpeg or use the Docker image.')
    try:
        result = subprocess.run(
            [executable, '-v', 'error', '-nostdin', '-i', str(video_path),
             '-map', '0:v:0', '-frames:v', '1',
             '-vf', 'scale=trunc(iw*sar+0.5):ih,setsar=1',
             '-f', 'image2pipe', '-c:v', 'png', 'pipe:1'],
            capture_output=True, timeout=30, check=False,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise RuntimeError(f'Cannot extract video preview from {video_path}: {exc}') from exc
    if result.returncode or not result.stdout:
        detail = result.stderr.decode('utf-8', errors='replace').strip()[-1500:]
        raise RuntimeError(f'Cannot extract video preview from {video_path}: {detail or "no video frame"}')
    return result.stdout


def contain_bounds(bounds, image_width: int, image_height: int) -> tuple[int, int, int, int]:
    """Scale proportionally, then center inside the PDF placeholder."""
    left, top, width, height = bounds
    scale = min(width / image_width, height / image_height)
    fitted_width = min(width, max(1, round(image_width * scale)))
    fitted_height = min(height, max(1, round(image_height * scale)))
    return (left + (width - fitted_width) // 2, top + (height - fitted_height) // 2,
            fitted_width, fitted_height)


def video_placeholder_rect(page, link_rect) -> pymupdf.Rect:
    """Find the smallest rectangular stroke enclosing a text-only TeX link.

    XeTeX can emit a link covering just the text inside an fbox. Recognize
    both PDF rectangles and fbox's four separate rules, but never expand to
    an unclosed set of lines. Work in unrotated drawing coordinates.
    """
    anchor = pymupdf.Rect(link_rect) * page.derotation_matrix
    candidates = [anchor]
    rectangles = []
    horizontal, vertical = [], []
    for drawing in page.get_drawings():
        if 's' not in drawing['type']:
            continue
        for item in drawing['items']:
            if item[0] == 're':
                rectangles.append(pymupdf.Rect(item[1]))
            elif item[0] == 'l':
                a, b = item[1:3]
                if abs(a.y - b.y) < 0.01:
                    horizontal.append((min(a.x, b.x), max(a.x, b.x), a.y))
                elif abs(a.x - b.x) < 0.01:
                    vertical.append((a.x, min(a.y, b.y), max(a.y, b.y)))
    # TeX's horizontal and vertical rule endpoints can differ by half a stroke.
    tolerance = 1.0
    # XeTeXLinkBox adds a 2pt safety margin. Prefer the actual frame when the
    # annotation's four edges are all this close, rather than enlarging it.
    link_margin = 2.5
    tops = [line for line in horizontal if line[2] <= anchor.y0 + link_margin
            and line[0] <= anchor.x0 + link_margin and line[1] >= anchor.x1 - link_margin]
    bottoms = [line for line in horizontal if line[2] >= anchor.y1 - link_margin
               and line[0] <= anchor.x0 + link_margin and line[1] >= anchor.x1 - link_margin]
    for x0, x1, y0 in tops:
        for bx0, bx1, y1 in bottoms:
            if abs(x0 - bx0) > tolerance or abs(x1 - bx1) > tolerance:
                continue
            def has_side(x):
                return any(abs(vx - x) <= tolerance and vy0 <= y0 + tolerance
                           and vy1 >= y1 - tolerance for vx, vy0, vy1 in vertical)
            if has_side(x0) and has_side(x1):
                rectangles.append(pymupdf.Rect(x0, y0, x1, y1))
    enclosed = [rect for rect in rectangles if not rect.is_empty and
                (rect.contains(anchor) or all(abs(a - b) <= link_margin for a, b in zip(rect, anchor)))]
    if enclosed:
        candidates = enclosed
    return min(candidates, key=lambda rect: rect.get_area()) * page.rotation_matrix


def render_page(page, dpi: int, video_rects: list[pymupdf.Rect]):
    """Erase movie placeholders on an in-memory copy, leaving the source PDF intact.

    Contained drawings/images and placeholder text are removed; larger vector
    panels and background images remain. Original links are read separately.
    """
    if not video_rects:
        return page.get_pixmap(dpi=dpi, colorspace=pymupdf.csRGB, alpha=False)
    with pymupdf.open() as clean:
        clean.insert_pdf(page.parent, from_page=page.number, to_page=page.number)
        target = clean[0]
        areas = []
        for rect in video_rects:
            # Half a point includes centered strokes and their anti-aliasing.
            area = pymupdf.Rect(rect) * page.derotation_matrix
            area = pymupdf.Rect(area.x0 - .5, area.y0 - .5, area.x1 + .5, area.y1 + .5)
            areas.append(area)
            target.add_redact_annot(area, fill=None, cross_out=False)
        # Image redaction would whiten parts of a photographic background.
        # Delete only images wholly inside placeholders, and only when every
        # occurrence of that image on this page is contained in a placeholder.
        for xref in {image[0] for image in target.get_images()}:
            boxes = target.get_image_rects(xref)
            if boxes and all(any(area.contains(box) for area in areas) for box in boxes):
                target.delete_image(xref)
        target.apply_redactions(images=0, graphics=1, text=0)
        return target.get_pixmap(dpi=dpi, colorspace=pymupdf.csRGB, alpha=False)


def make_invisible_hotspot(shape) -> None:
    """Use a transparent *filled* shape, retaining its entire clickable area."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    color = shape._element.spPr.solidFill[0]
    alpha = OxmlElement('a:alpha')
    # Keep a tiny amount of opacity: PowerPoint versions differ in whether a
    # shape with a fully transparent fill participates in hit testing.
    alpha.set('val', '1000')
    color.append(alpha)
    shape.line.fill.background()
    shape._element.spPr.append(OxmlElement('a:effectLst'))


def hotspot_bounds(rect, page_rect, placement) -> tuple[int, int, int, int] | None:
    if rect is None:
        return None
    rect = pymupdf.Rect(rect)
    if not all(math.isfinite(v) for v in rect) or rect.is_empty or rect.is_infinite:
        return None
    rect &= page_rect
    if rect.is_empty:
        return None
    left, top, width, height = placement
    sx, sy = width / page_rect.width, height / page_rect.height
    x0 = left + round((rect.x0 - page_rect.x0) * sx)
    y0 = top + round((rect.y0 - page_rect.y0) * sy)
    x1 = left + round((rect.x1 - page_rect.x0) * sx)
    y1 = top + round((rect.y1 - page_rect.y0) * sy)
    if x1 <= x0 or y1 <= y0:
        return None
    return x0, y0, x1 - x0, y1 - y0


def save_presentation(prs, path: Path) -> None:
    """Replace output only after a complete archive has been written."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix='.pptx', delete=False) as temp:
        temporary_path = Path(temp.name)
    try:
        prs.save(str(temporary_path))
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def set_slide_transition(slide, page) -> bool:
    """Preserve explicit PDF effects/timing without guessing overlay boundaries."""
    pdf = page.parent
    _, style = pdf.xref_get_key(page.xref, 'Trans/S')
    effect = {'/Fade': 'fade', '/Dissolve': 'dissolve'}.get(style)
    if style not in ('null', '/R', '/Fade', '/Dissolve'):
        print(f'WARNING: page {page.number + 1}: unsupported PDF transition {style}; '
              'using an ordinary slide change. Use transfade or transdissolve.', file=sys.stderr)

    def milliseconds(key, default=None):
        kind, value = pdf.xref_get_key(page.xref, key)
        if kind == 'null':
            return default
        if kind in ('int', 'float'):
            seconds = float(value)
            if math.isfinite(seconds) and 0 <= seconds <= 4294967.295:
                return round(seconds * 1000)
        print(f'WARNING: page {page.number + 1}: invalid {key} duration; ignored.', file=sys.stderr)
        return default

    duration = milliseconds('Trans/D', 1000) if effect else None
    advance = milliseconds('Dur')
    if effect is None and advance is None:
        return False
    transition = OxmlElement('p:transition')
    transition.set('advClick', '1')
    if effect:
        # Office 2010+ reads the exact milliseconds; older clients use spd.
        transition = parse_xml(
            '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" '
            'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            'mc:Ignorable="p14" advClick="1"/>')
        transition.set('{http://schemas.microsoft.com/office/powerpoint/2010/main}dur', str(duration))
        transition.set('spd', 'fast' if duration <= 500 else 'med' if duration <= 1000 else 'slow')
        transition.append(OxmlElement('p:' + effect))
    if advance is not None:
        transition.set('advTm', str(advance))
    slide._element.insert_element_before(transition, 'p:timing', 'p:extLst')
    return True


def convert_pdf_to_pptx(
    pdf_path: str | Path,
    pptx_path: str | Path,
    dpi: int = 600,
    verbose: bool = False,
    strict_links: bool = False,
    notes_path: str | Path | None = None,
    auto_notes: bool = True,
    strict_notes: bool = False,
    notes_dir: str | Path | None = None,
    preserve_transitions: bool = True,
    notes_pdf_path: str | Path | None = None,
) -> LinkSummary:
    pdf_path = Path(pdf_path).resolve()
    pptx_path = Path(pptx_path).resolve()
    if not pdf_path.is_file():
        raise FileNotFoundError(pdf_path)
    if pdf_path == pptx_path or (pptx_path.exists() and pdf_path.samefile(pptx_path)):
        raise ValueError('Input PDF and output PPTX must be different files.')
    if pptx_path.suffix.lower() != '.pptx':
        raise ValueError('Output must use the .pptx extension.')
    if type(dpi) is not int or dpi <= 0:
        raise ValueError('DPI must be a positive integer.')
    if notes_path is not None and notes_pdf_path is not None:
        raise ValueError('Choose either native PDF notes or legacy JSON notes.')
    automatic_notes = notes_path is None and notes_pdf_path is None and auto_notes
    if automatic_notes:
        directory = Path(notes_dir).resolve() if notes_dir is not None else pdf_path.parent
        if not directory.is_dir():
            raise ValueError(f'Notes directory does not exist: {directory}')
        native = directory / (pdf_path.stem + '.notes.pdf')
        legacy = directory / (pdf_path.stem + '.notes.json')
        if native.is_file():
            with pymupdf.open(native) as rendered:
                has_notes = pdf_info_string(rendered, 'CampusHasNotes')
            if has_notes == 'false' and legacy.is_file():
                notes_path = legacy
            else:
                notes_pdf_path = native
        else:
            if legacy.is_file():
                notes_path = legacy

    with pymupdf.open(pdf_path) as pdf:
        if not pdf.is_pdf:
            raise ValueError('Input must be a PDF document.')
        if pdf.needs_pass:
            raise ValueError('Input PDF is encrypted. Decrypt it before conversion.')
        if not pdf.page_count:
            raise ValueError('PDF contains no pages.')
        if pdf_info_string(pdf, 'CampusNotesLayout') == 'right':
            raise ValueError('This PDF includes a notes screen. Run make to export slides only to PPTX.')
        notes = (load_native_notes(notes_pdf_path, pdf, pdf_path) if notes_pdf_path is not None
                 else load_notes(notes_path, pdf.page_count, strict=strict_notes))
        if notes_path is None and notes_pdf_path is None:
            print('Notes: no notes file selected or found; speaker notes will be empty.')
        else:
            print(f'Notes: {notes_pdf_path or notes_path}')
        try:
            named_destinations = pdf.resolve_names()
        except RuntimeError as exc:
            print(f'WARNING: cannot read named destinations: {exc}', file=sys.stderr)
            named_destinations = {}

        prs = Presentation()
        first_rect = pdf[0].rect
        prs.slide_width = Inches(13.333)
        prs.slide_height = round(prs.slide_width * first_rect.height / first_rect.width)
        slides = [prs.slides.add_slide(prs.slide_layouts[6]) for _ in pdf]
        campus = campus_metadata(pdf)
        set_document_metadata(prs, pdf, campus)
        section_count = set_presentation_sections(prs, pdf, campus)
        print(f'Sections: {section_count} | Metadata: title, author, subject, keywords; {len(campus)} template fields')
        for slide, note in zip(slides, notes):
            set_slide_notes(slide, note)
        summary = LinkSummary()
        transition_count = 0
        posters: dict[Path, bytes] = {}
        print(f'Input: {pdf_path}\nPages: {len(pdf)} | DPI: {dpi}')

        for page_index, page in enumerate(pdf):
            slide = slides[page_index]
            # A PPTX has one slide size. Fit mixed-size PDF pages without
            # stretching, and use exactly the same transform for link areas.
            scale = min(prs.slide_width / page.rect.width, prs.slide_height / page.rect.height)
            width, height = round(page.rect.width * scale), round(page.rect.height * scale)
            left, top = (prs.slide_width - width) // 2, (prs.slide_height - height) // 2
            placement = (left, top, width, height)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
            links = page.get_links()
            videos = {}
            for index, link in enumerate(links):
                video = local_video_for_link(link, pdf_path)
                if video is None or hotspot_bounds(link.get('from'), page.rect, placement) is None:
                    continue
                video_path, mime_type = video
                if video_path not in posters:
                    posters[video_path] = video_poster(video_path)
                rect = video_placeholder_rect(page, link['from']) & page.rect
                videos[index] = (video_path, mime_type, rect)
            pixmap = render_page(page, dpi, [video[2] for video in videos.values()])
            with BytesIO(pixmap.tobytes('png')) as image:
                slide.shapes.add_picture(image, *placement)
            del pixmap

            # MuPDF omits some unsupported actions (e.g. JavaScript) entirely.
            # Include those in the skip count instead of claiming full coverage.
            annotation_count = sum(kind == pymupdf.PDF_ANNOT_LINK for _, kind, _ in page.annot_xrefs())
            missing = max(0, annotation_count - len(links))
            summary.skipped += missing
            if missing and verbose:
                print(f'  page {page_index + 1}: {missing} unreadable link annotation(s)')
            for link_index, link in enumerate(links):
                target_page = resolve_internal_link(link, page_index, len(pdf), named_destinations)
                uri = None if target_page is not None else resolve_external_link(link, pdf_path, pptx_path)
                # PyMuPDF link rectangles already follow the page's displayed
                # rotation; applying rotation_matrix again would misplace them.
                bounds = hotspot_bounds(link.get('from'), page.rect, placement)
                if bounds is None or (target_page is None and uri is None):
                    summary.skipped += 1
                    if verbose:
                        print(f'  [SKIPPED] page {page_index + 1}: {link!r}')
                    continue
                video = videos.get(link_index)
                if video is not None:
                    video_path, mime_type, rect = video
                    poster = posters[video_path]
                    image = pymupdf.Pixmap(poster)
                    bounds = contain_bounds(hotspot_bounds(rect, page.rect, placement), image.width, image.height)
                    with BytesIO(poster) as preview:
                        movie = slide.shapes.add_movie(
                            str(video_path), *bounds, poster_frame_image=preview, mime_type=mime_type
                        )
                    movie.name = f'Embedded video: {video_path.name}'
                    summary.embedded += 1
                    if verbose:
                        print(f'  page {page_index + 1} -> embedded {video_path}')
                    continue
                hotspot = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *bounds)
                make_invisible_hotspot(hotspot)
                if target_page is not None:
                    hotspot.click_action.target_slide = slides[target_page]
                    hotspot.name = f'Link to slide {target_page + 1}'
                    summary.internal += 1
                else:
                    hotspot.click_action.hyperlink.address = uri
                    hotspot.name = f'Link to {uri}'
                    summary.external += 1
                if verbose:
                    destination = f'slide {target_page + 1}' if target_page is not None else uri
                    print(f'  page {page_index + 1} -> {destination}')
            print(f'Converted page {page_index + 1}/{len(pdf)}')
            if preserve_transitions:
                transition_count += set_slide_transition(slide, page)

        if summary.skipped:
            message = f'{summary.skipped} link(s) could not be preserved; use --verbose for details.'
            if strict_links:
                raise ValueError(message + ' Output was not saved.')
            print(f'WARNING: {message}', file=sys.stderr)
        set_export_properties(prs, summary.embedded)
        save_presentation(prs, pptx_path)
        print(f'Saved: {pptx_path}')
        print(f'Transitions: {transition_count} slides with explicit PDF effects/timing')
        print(
            f'Links: {summary.internal} internal, {summary.external} external, '
            f'{summary.embedded} embedded video, {summary.skipped} skipped'
        )
        return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('input_pdf', type=Path, help='Input Beamer/PDF file')
    parser.add_argument('output_pptx', type=Path, nargs='?', help='Output file (default: input name with .pptx)')
    parser.add_argument('--dpi', type=int, default=600, help='Rendering DPI (default: 600)')
    parser.add_argument('--notes', type=Path, help='JSON speaker notes, one entry per final PDF page')
    parser.add_argument('--notes-pdf', type=Path, help='Native Campus second-screen notes PDF generated from TeX')
    parser.add_argument('--notes-dir', type=Path, help='Auto-load native PDF or legacy JSON notes from this directory (default: PDF directory)')
    parser.add_argument('--no-notes', action='store_true', help='Do not auto-load PDF or JSON notes')
    parser.add_argument('-v', '--verbose', action='store_true', help='Show link destinations and skipped links')
    parser.add_argument('--strict-links', action='store_true', help='Do not save if any link is unsupported or invalid')
    parser.add_argument('--strict-notes', action='store_true', help='If notes are loaded, require exactly one entry per PDF page')
    parser.add_argument('--no-transitions', action='store_true', help='Ignore PDF transition effects and automatic page durations')
    args = parser.parse_args()
    try:
        convert_pdf_to_pptx(
            args.input_pdf, args.output_pptx or args.input_pdf.with_suffix('.pptx'),
            args.dpi, args.verbose, args.strict_links,
            args.notes, auto_notes=not args.no_notes, strict_notes=args.strict_notes,
            notes_dir=args.notes_dir,
            preserve_transitions=not args.no_transitions,
            notes_pdf_path=args.notes_pdf,
        )
    except (OSError, RuntimeError, ValueError) as exc:
        parser.exit(1, f'Error: {exc}\n')


if __name__ == '__main__':
    main()
