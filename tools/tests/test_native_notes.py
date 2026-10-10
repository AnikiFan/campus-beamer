# Copyright (C) 2026 Xiao Fan <xiaofan140@gmail.com>
# SPDX-License-Identifier: GPL-3.0-or-later
"""Native notes preserve empty stages, privacy, and the existing JSON import API."""
from contextlib import redirect_stdout
import hashlib
from io import StringIO
import json
from pathlib import Path
import sys
import tempfile
import unittest

import pymupdf
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from images_to_ppt import convert_pdf_to_pptx


def info(pdf, key, value):
    kind, reference = pdf.xref_get_key(-1, 'Info')
    if kind != 'xref':
        pdf.set_metadata({'title': 'Notes fixture'})
        _, reference = pdf.xref_get_key(-1, 'Info')
    pdf.xref_set_key(int(reference.split()[0]), key, '(' + value + ')')


class NativeNotesTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.pdf = self.root / 'slides.pdf'
        self.notes = self.root / 'slides.notes.pdf'
        self.pptx = self.root / 'slides.pptx'
        with pymupdf.open() as pdf:
            for index in range(3):
                pdf.new_page(width=320, height=180).insert_text((20, 50), f'Visible {index}')
            info(pdf, 'CampusNotesLayout', 'none')
            pdf.save(self.pdf)

    def note_pdf(self, texts=('First note', '', 'Last note'), *, width=640, layout='right', fingerprint=None, has_notes='true'):
        with pymupdf.open() as pdf:
            for index, text in enumerate(texts):
                page = pdf.new_page(width=width, height=180)
                page.insert_text((20, 50), f'Audience content {index}')
                if text:
                    page.insert_text((width / 2 + 20, 50), text)
            info(pdf, 'CampusNotesLayout', layout)
            info(pdf, 'CampusHasNotes', has_notes)
            info(pdf, 'CampusSlidesSHA256', fingerprint or hashlib.sha256(self.pdf.read_bytes()).hexdigest())
            pdf.save(self.notes)

    def convert(self, **options):
        with redirect_stdout(StringIO()):
            convert_pdf_to_pptx(self.pdf, self.pptx, dpi=36, strict_notes=True, **options)
        return Presentation(self.pptx)

    def test_native_notes_take_precedence_over_legacy_and_keep_blank_stages(self):
        self.note_pdf()
        self.pdf.with_suffix('.notes.json').write_text(json.dumps(['Old notes'] * 3))
        ppt = self.convert()
        self.assertEqual([s.notes_slide.notes_text_frame.text for s in ppt.slides],
                         ['First note', '', 'Last note'])
        for slide in ppt.slides:
            self.assertNotIn('Audience', slide.notes_slide.notes_text_frame.text)
        self.assertAlmostEqual(ppt.slide_width / ppt.slide_height, 320 / 180, places=5)

    def test_explicit_legacy_json_still_overrides_auto_discovery(self):
        self.note_pdf()
        legacy = self.root / 'legacy.json'
        legacy.write_text(json.dumps(['Explicit'] * 3))
        ppt = self.convert(notes_path=legacy)
        self.assertEqual(ppt.slides[0].notes_slide.notes_text_frame.text, 'Explicit')

    def test_old_talk_without_native_commands_retains_legacy_notes(self):
        self.note_pdf(('', '', ''), has_notes='false')
        self.pdf.with_suffix('.notes.json').write_text(json.dumps(['Existing talk notes'] * 3))
        ppt = self.convert()
        self.assertEqual(ppt.slides[0].notes_slide.notes_text_frame.text, 'Existing talk notes')

    def test_explicit_empty_native_commands_do_not_restore_old_json(self):
        self.note_pdf(('', '', ''))
        self.pdf.with_suffix('.notes.json').write_text(json.dumps(['Stale text'] * 3))
        ppt = self.convert()
        self.assertEqual([s.notes_slide.notes_text_frame.text for s in ppt.slides], ['', '', ''])

    def test_count_mismatch_preserves_the_previous_pptx(self):
        self.note_pdf(('Only one',))
        self.pptx.write_bytes(b'Previous valid export')
        with self.assertRaisesRegex(ValueError, 'different page counts'):
            self.convert()
        self.assertEqual(self.pptx.read_bytes(), b'Previous valid export')

    def test_stale_notes_are_rejected_even_with_equal_page_counts(self):
        self.note_pdf(fingerprint='0' * 64)
        with self.assertRaisesRegex(ValueError, 'does not match'):
            self.convert()
        self.assertFalse(self.pptx.exists())

    def test_unmarked_or_wrong_geometry_note_pdfs_are_rejected(self):
        for layout, width, diagnostic in (('none', 640, 'second-screen'), ('right', 320, 'geometry')):
            with self.subTest(layout=layout):
                self.note_pdf(layout=layout, width=width)
                with self.assertRaisesRegex(ValueError, diagnostic):
                    self.convert()

    def test_disabling_auto_notes_keeps_notes_empty(self):
        self.note_pdf()
        ppt = self.convert(auto_notes=False)
        self.assertEqual([s.notes_slide.notes_text_frame.text for s in ppt.slides], ['', '', ''])

    def test_a_second_screen_pdf_is_not_exported_as_a_wide_slide(self):
        with pymupdf.open(self.pdf) as pdf:
            info(pdf, 'CampusNotesLayout', 'right')
            pdf.saveIncr()
        with self.assertRaisesRegex(ValueError, 'notes screen'):
            self.convert(auto_notes=False)
        self.assertFalse(self.pptx.exists())

    def test_explicit_pdf_and_json_notes_cannot_be_mixed(self):
        self.note_pdf()
        with self.assertRaisesRegex(ValueError, 'either native'):
            self.convert(notes_path=self.root / 'unused.json', notes_pdf_path=self.notes)


if __name__ == '__main__':
    unittest.main()
