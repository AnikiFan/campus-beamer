# Copyright (C) 2026 Xiao Fan <xiaofan140@gmail.com>
# Part of Campus Beamer. See LICENSE for the GNU GPL v3 notice.
"""Native slide transitions preserve explicit PDF intent and export contracts."""
from contextlib import redirect_stderr, redirect_stdout
from io import BytesIO, StringIO
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

import pymupdf
from pptx import Presentation
from pptx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from images_to_ppt import convert_pdf_to_pptx, set_slide_transition

P14_DURATION = '{http://schemas.microsoft.com/office/powerpoint/2010/main}dur'


class TransitionTests(unittest.TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.root = Path(folder.name)
        self.pdf = self.root / 'stages.pdf'
        self.pptx = self.root / 'stages.pptx'

    def document(self, count):
        document = pymupdf.open()
        self.addCleanup(document.close)
        for index in range(count):
            page = document.new_page(width=320, height=180)
            page.insert_text((20, 60), f'Stage {index + 1}')
        return document

    def convert(self, **options):
        with redirect_stdout(StringIO()), redirect_stderr(StringIO()) as warnings:
            convert_pdf_to_pptx(self.pdf, self.pptx, dpi=36, **options)
        return Presentation(self.pptx), warnings.getvalue()

    def test_direct_and_indirect_effects_keep_notes_links_and_sections(self):
        pdf = self.document(3)
        pdf.xref_set_key(pdf[1].xref, 'Trans', '<< /S /Fade /D 0.2 >>')
        indirect = pdf.get_new_xref()
        pdf.update_object(indirect, '<< /S /Dissolve /D 0.45 >>')
        pdf.xref_set_key(pdf[2].xref, 'Trans', f'{indirect} 0 R')
        pdf[0].insert_link({'kind': pymupdf.LINK_GOTO, 'from': pymupdf.Rect(10, 10, 90, 25), 'page': 2})
        pdf[1].insert_link({'kind': pymupdf.LINK_URI, 'from': pymupdf.Rect(10, 10, 90, 25), 'uri': 'https://example.org/'})
        pdf.set_toc([[1, 'Steps', 1]])
        pdf.save(self.pdf)
        self.pdf.with_suffix('.notes.json').write_text(json.dumps(
            {'slides': [{'notes': f'Explain stage {i + 1}'} for i in range(3)]}))
        original = self.pdf.read_bytes()
        ppt, warnings = self.convert(strict_notes=True, strict_links=True)
        self.assertEqual(warnings, '')
        self.assertEqual(len(ppt.slides), 3)
        self.assertIsNone(ppt.slides[0]._element.find(qn('p:transition')))
        for index, effect, duration in ((1, 'fade', '200'), (2, 'dissolve', '450')):
            transition = ppt.slides[index]._element.find(qn('p:transition'))
            self.assertIsNotNone(transition.find(qn('p:' + effect)))
            self.assertEqual(transition.get(P14_DURATION), duration)
            self.assertEqual(transition.get('advClick'), '1')
            self.assertIsNone(transition.get('advTm'))
        self.assertEqual([s.notes_slide.notes_text_frame.text for s in ppt.slides],
                         [f'Explain stage {i + 1}' for i in range(3)])
        self.assertEqual(ppt.slides[0].shapes[1].click_action.target_slide.slide_id,
                         ppt.slides[2].slide_id)
        self.assertEqual(ppt.slides[1].shapes[1].click_action.hyperlink.address, 'https://example.org/')
        self.assertIn('Steps', ppt._element.xml)
        self.assertEqual(self.pdf.read_bytes(), original)

    def test_automatic_advance_requires_explicit_page_duration(self):
        pdf = self.document(3)
        pdf.xref_set_key(pdf[0].xref, 'Trans', '<< /S /R >>')
        pdf.xref_set_key(pdf[1].xref, 'Dur', '1.25')
        pdf.xref_set_key(pdf[2].xref, 'Trans', '<< /S /Fade >>')
        pdf.xref_set_key(pdf[2].xref, 'Dur', '0')
        pdf.save(self.pdf)
        ppt, warnings = self.convert()
        self.assertEqual(warnings, '')
        self.assertIsNone(ppt.slides[0]._element.find(qn('p:transition')))
        self.assertEqual(ppt.slides[1]._element.find(qn('p:transition')).get('advTm'), '1250')
        transition = ppt.slides[2]._element.find(qn('p:transition'))
        self.assertEqual(transition.get('advTm'), '0')
        self.assertEqual(transition.get(P14_DURATION), '1000')

    def test_unsupported_effect_warns_and_retains_the_page(self):
        pdf = self.document(1)
        pdf.xref_set_key(pdf[0].xref, 'Trans', '<< /S /Fly >>')
        pdf.save(self.pdf)
        ppt, warnings = self.convert()
        self.assertIn('unsupported PDF transition', warnings)
        self.assertEqual(len(ppt.slides), 1)
        self.assertIsNone(ppt.slides[0]._element.find(qn('p:transition')))

    def test_invalid_durations_do_not_create_unintended_timers(self):
        pdf = self.document(3)
        for page, value in zip(pdf, ('-1', 'true', '4294968')):
            pdf.xref_set_key(page.xref, 'Dur', value)
        pdf.xref_set_key(pdf[0].xref, 'Trans', '<< /S /Fade /D -1 >>')
        pdf.save(self.pdf)
        ppt, warnings = self.convert()
        self.assertEqual(warnings.count('invalid'), 4)
        for slide in ppt.slides:
            transition = slide._element.find(qn('p:transition'))
            if transition is not None:
                self.assertIsNone(transition.get('advTm'))
        self.assertEqual(ppt.slides[0]._element.find(qn('p:transition')).get(P14_DURATION), '1000')

    def test_cli_can_disable_effects_and_timers(self):
        pdf = self.document(1)
        pdf.xref_set_key(pdf[0].xref, 'Trans', '<< /S /Fade /D 0.2 >>')
        pdf.xref_set_key(pdf[0].xref, 'Dur', '1')
        pdf.save(self.pdf)
        script = Path(__file__).resolve().parents[1] / 'images_to_ppt.py'
        result = subprocess.run([sys.executable, str(script), str(self.pdf), str(self.pptx),
                                 '--dpi', '36', '--no-notes', '--no-transitions'],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIsNone(Presentation(self.pptx).slides[0]._element.find(qn('p:transition')))

    def test_transition_precedes_existing_video_timing(self):
        pdf = self.document(1)
        page = pdf[0]
        pdf.xref_set_key(page.xref, 'Trans', '<< /S /Fade /D 0.2 >>')
        ppt = Presentation()
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        video = Path(__file__).resolve().parents[2] / 'assets/presentation_demo.mp4'
        slide.shapes.add_movie(str(video), 0, 0, 1000000, 500000,
                              poster_frame_image=BytesIO(page.get_pixmap().tobytes('png')),
                              mime_type='video/mp4')
        timing = slide._element.find(qn('p:timing'))
        original = timing.xml
        set_slide_transition(slide, page)
        children = list(slide._element)
        self.assertLess(children.index(slide._element.find(qn('p:transition'))), children.index(timing))
        self.assertEqual(timing.xml, original)


if __name__ == '__main__':
    unittest.main()
